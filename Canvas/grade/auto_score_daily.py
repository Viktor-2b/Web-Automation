import pandas as pd
from openpyxl import load_workbook
from decimal import Decimal, ROUND_HALF_UP  # 引入高精度计算库


def calculate_and_update_daily():
    # 1. 读取平时分文件
    try:
        df = pd.read_excel('平时分.xlsx')
    except FileNotFoundError:
        print("❌ 找不到 平时分.xlsx 文件")
        return

    # 一共8项指标
    score_columns = ['作业1', '作业2', '作业3', '作业4', '作业5', '作业6', '课堂小测', '出勤']

    # 检查是否有列名没对上
    missing_cols = [col for col in score_columns if col not in df.columns]
    if missing_cols:
        print(f"⚠️ 警告：平时分.xlsx 中找不到以下列：{missing_cols}，将只计算存在的列。")
        score_columns = [col for col in score_columns if col in df.columns]

    # 使用 Decimal 进行绝对精确的求平均与两位小数四舍五入
    def calculate_true_mean_2_decimals(row):
        # 统一转成 Decimal 防精度丢失
        total = sum(Decimal(str(x)) for x in row)
        avg = total / Decimal(len(row))
        # 强制四舍五入到两位小数 (0.01)，并转回 float
        return float(avg.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP))

    # 应用高精度计算
    df['最终平时分'] = df[score_columns].apply(calculate_true_mean_2_decimals, axis=1)

    # 建立 姓名 -> 平时分 的映射字典
    df['姓名'] = df['姓名'].astype(str).str.strip()
    score_map = dict(zip(df['姓名'], df['最终平时分']))

    # 2. 写入 总评.xlsx
    target_file = '总评.xlsx'
    try:
        wb = load_workbook(target_file)
    except FileNotFoundError:
        print(f"❌ 找不到文件: {target_file}")
        return

    ws = wb.active

    # 获取表头索引，防空格干扰
    headers = {}
    for idx, cell in enumerate(ws[1]):
        if cell.value is not None:
            headers[str(cell.value).strip()] = idx

    col_name = headers.get('姓名')
    col_daily = headers.get('平时分')  # 请确保你的总评表里这一列叫“平时分”

    if col_name is None:
        print("❌ 写入失败：总评.xlsx 中未找到 '姓名' 列。")
        return
    if col_daily is None:
        print("❌ 写入失败：总评.xlsx 中未找到 '平时分' 列，请检查表头是否一致。")
        return

    update_count = 0
    # 遍历总评表并填充分数
    for row in ws.iter_rows(min_row=2):
        raw_name = row[col_name].value
        if raw_name:
            name = str(raw_name).strip()
            if name in score_map:
                # 写入数据并强制规定 Excel 渲染为两位小数
                cell_daily = row[col_daily]
                cell_daily.value = score_map[name]
                cell_daily.number_format = '0.00'  # 让 Excel 显示如 97.50
                update_count += 1

    wb.save(target_file)
    print(f"✅ 平时分高精度计算完毕！已将 {update_count} 人的成绩(保留两位小数)写入 {target_file} 的 '平时分' 列。")


if __name__ == '__main__':
    calculate_and_update_daily()