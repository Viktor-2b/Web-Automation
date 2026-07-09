import numpy as np
import matplotlib.pyplot as plt
from openpyxl import load_workbook
from scipy.stats import gaussian_kde
from decimal import Decimal, ROUND_HALF_UP, InvalidOperation

# 解决图表中的中文显示问题 (兼容 Windows / Mac)
plt.rcParams['font.sans-serif'] = ['SimHei', 'Microsoft YaHei', 'Arial Unicode MS']
plt.rcParams['axes.unicode_minus'] = False


def calculate_final_score():
    filename = '总评.xlsx'
    try:
        wb = load_workbook(filename)
    except FileNotFoundError:
        print(f"❌ 找不到文件: {filename}")
        return

    ws = wb.active

    # 1. 动态获取表头
    headers = {}
    for idx, cell in enumerate(ws[1]):
        if cell.value is not None:
            headers[str(cell.value).strip()] = idx

    col_daily = headers.get('平时分')
    col_major_project = headers.get('大作业')
    col_final_exam = headers.get('期末考试') or headers.get('期末')
    col_total = headers.get('总评成绩') or headers.get('总评') or headers.get('总分')

    col_grade = headers.get('等级')
    if col_grade is None:
        col_grade = len(headers)
        ws.cell(row=1, column=col_grade + 1, value='等级')
        print("💡 表格中未找到[等级]列，已自动在末尾创建。")

    if None in [col_daily, col_major_project, col_final_exam, col_total]:
        print(f"❌ 缺少必要的列头！请检查表头。")
        return

    update_count = 0
    valid_final_scores = []

    # 3. 遍历每一行计算最终总评
    for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
        val_p = row[col_daily].value
        val_d = row[col_major_project].value
        val_q = row[col_final_exam].value

        if val_p is None and val_d is None and val_q is None:
            continue

        try:
            score_p = Decimal(str(val_p)) if val_p is not None else Decimal('0.0')
            score_d = Decimal(str(val_d)) if val_d is not None else Decimal('0.0')
            score_q = Decimal(str(val_q)) if val_q is not None else Decimal('0.0')
        except (ValueError, TypeError, InvalidOperation):
            continue

        # 乘以绝对精确的权重
        raw_final = score_p * Decimal('0.4') + score_d * Decimal('0.3') + score_q * Decimal('0.3')

        # 精确保留两位小数 (0.01)，并执行标准的逢五进一
        final_score_decimal = raw_final.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

        # 转回 float 类型，方便下方评级（如 89.99 < 90，严格卡档）和画图
        final_score = float(final_score_decimal)

        # 兜底保护，最高100.0
        final_score = min(final_score, 100.0)

        if final_score >= 95:
            grade = 'A+'
        elif final_score >= 90:
            grade = 'A'
        elif final_score >= 85:
            grade = 'A-'
        elif final_score >= 80:
            grade = 'B+'
        elif final_score >= 75:
            grade = 'B'
        elif final_score >= 70:
            grade = 'B-'
        elif final_score >= 67:
            grade = 'C+'
        elif final_score >= 65:
            grade = 'C'
        elif final_score >= 62:
            grade = 'C-'
        elif final_score >= 60:
            grade = 'D'
        else:
            grade = 'F'

        # 获取总评单元格对象
        cell_total = row[col_total]
        # 填入数值
        cell_total.value = final_score
        # 强制 Excel 将该单元格渲染为严格的两位小数格式
        cell_total.number_format = '0.00'

        ws.cell(row=row_idx, column=col_grade + 1, value=grade)
        update_count += 1

        if final_score > 0:
            valid_final_scores.append(final_score)

    wb.save(filename)
    print(f"✅ 4-3-3 总评及等级计算完毕！成功更新了 {update_count} 人的最终成绩。")

    # ================= 统计优秀率 =================
    if valid_final_scores:
        total_students = len(valid_final_scores)
        a_count = len([s for s in valid_final_scores if s >= 90])
        a_minus_count = len([s for s in valid_final_scores if s >= 85])

        a_rate = (a_count / total_students) * 100
        a_minus_rate = (a_minus_count / total_students) * 100

        print("\n" + "=" * 50)
        print("📊 成绩分布统计结果：")
        print(f"   - 录入总人数：{total_students}")
        print(f"   - A 及以上 (>=90分)：{a_count} 人，占比 {a_rate:.1f}%")
        print(f"   - A- 及以上 (>=85分)：{a_minus_count} 人，占比 {a_minus_rate:.1f}%")

        if a_rate <= 30.5:
            print("   👉 结论：优秀率(A及以上)合规！教务处可安全过审。")
        else:
            print("   ⚠️ 警告：优秀率(A及以上)超标！")
        print("=" * 50 + "\n")

    # ================= 画图环节 =================
    if valid_final_scores:
        print("正在生成平滑的成绩分布双曲线图...")

        # 【核心】：计算正态拟合曲线时，依然排除掉极端的低分（<60分），保证红线不会被那个29分的人扯歪
        normal_fit_scores = [s for s in valid_final_scores if s >= 60]
        mean_score = np.mean(normal_fit_scores)
        std_score = np.std(normal_fit_scores)

        plt.figure(figsize=(10, 6))

        # 构建 X 轴区间 (横跨 25 到 105 分，为了涵盖所有可能的低分)
        x = np.linspace(25, 105, 500)

        # 1. 绘制【真实成绩分布密度曲线】(KDE) 替代直方图
        kde = gaussian_kde(valid_final_scores, bw_method=0.4)
        y_real = kde(x)

        plt.plot(x, y_real, color='royalblue', linewidth=2.5, label='今年真实成绩曲线')
        plt.fill_between(x, y_real, alpha=0.3, color='royalblue')

        # 2. 绘制【大部队正态拟合曲线】(仅使用正常及格样本算出的均值和方差)
        p = (1 / (std_score * np.sqrt(2 * np.pi))) * np.exp(-0.5 * ((x - mean_score) / std_score) ** 2)
        plt.plot(x, p, color='crimson', linestyle='--', linewidth=2,
                 label=f'大部队正态拟合 ($\\mu={mean_score:.1f}, \\sigma={std_score:.1f}$)')

        # 添加分界线
        plt.axvline(85, color='green', linestyle=':', linewidth=1.5, label='A- 分界线 (85分)')
        plt.axvline(90, color='orange', linestyle=':', linewidth=1.5, label='A 分界线 (90分)')

        plt.title('计算机网络 总评成绩真实分布与正态拟合', fontsize=16)
        plt.xlabel('最终总评分数', fontsize=12)
        plt.ylabel('分布密度 (概率)', fontsize=12)
        plt.legend(fontsize=11, loc='upper left')
        plt.grid(axis='y', alpha=0.3)

        # 调整横坐标刻度，每5分显示一次
        plt.xticks(np.arange(25, 105, 5))

        plt.savefig('今年成绩分布双曲线图.png', dpi=300, bbox_inches='tight')
        print("📈 图表已保存为 '今年成绩分布双曲线图.png'，正在打开显示界面...")
        plt.show()


if __name__ == '__main__':
    calculate_final_score()