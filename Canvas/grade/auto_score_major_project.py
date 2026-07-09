import re
from openpyxl import load_workbook
from decimal import Decimal, ROUND_HALF_UP  # 新增：用于绝对精确的两位小数计算


def process_all_group_scores():
    sources = [
        '分组评分表 - 刘立伟.xlsx',
        '分组评分表 - 韩天柱.xlsx',
        '分组评分表 - 阮娜.xlsx'
    ]

    student_scores = {}

    # === 修复后的红字判断逻辑 ===
    def is_red_font(cell):
        if not cell.font or not cell.font.color:
            return False
        color = cell.font.color
        if color.type == 'rgb' and color.rgb:
            rgb_str = str(color.rgb).upper()
            # 黑色是 FF000000, 红色是 FFFF0000。必须以 FF0000 结尾才能避开黑色！
            if rgb_str.endswith('FF0000'):
                return True
        elif color.type == 'indexed' and color.indexed == 2:
            # 兼容有些旧版 Excel 中标准红色的 index 是 2
            return True
        return False

    # === 修复后的黄底判断逻辑 ===
    def is_yellow_fill(cell):
        if not cell.fill or not cell.fill.fgColor:
            return False
        color = cell.fill.fgColor
        if color.type == 'rgb' and color.rgb:
            rgb_str = str(color.rgb).upper()
            # 黄色通常是 FFFFFF00
            if rgb_str.endswith('FFFF00'):
                return True
        return False

    for src in sources:
        print("\n" + "=" * 50)
        print(f"📂 开始解析文件: {src}")
        print("=" * 50)

        try:
            wb = load_workbook(src, data_only=True)
            ws = wb.active
        except FileNotFoundError:
            print(f"❌ 找不到源文件: {src}，跳过此文件。")
            continue

        headers = {}
        for idx, cell in enumerate(ws[1]):
            if cell.value is not None:
                headers[str(cell.value).strip()] = idx

        col_group = headers.get('组号')
        col_total = headers.get('总分（100）')
        member_cols = [headers.get(f'成员{i}') for i in range(1, 5) if headers.get(f'成员{i}') is not None]

        if col_group is None or col_total is None:
            print(f"❌ 文件 {src} 缺少 '组号' 或 '总分（100）' 列，跳过解析。")
            continue

        for row in ws.iter_rows(min_row=2):
            raw_group_id = row[col_group].value
            if raw_group_id is None:
                continue

            group_id = int(raw_group_id) if isinstance(raw_group_id,
                                                       float) and raw_group_id.is_integer() else raw_group_id

            try:
                group_score = float(row[col_total].value)
            except (ValueError, TypeError):
                continue

            print(f"\n▶ 第 {group_id} 组 | 组总分: {group_score}")

            for m_col in member_cols:
                cell = row[m_col]
                raw_val = cell.value

                if raw_val:
                    raw_str = str(raw_val).strip()
                    score_match = re.search(r'[(（](\d+(?:\.\d+)?)[)）]', raw_str)
                    clean_name = re.sub(r'[(（].*?[)）]', '', raw_str).strip()

                    detail_log = ""

                    if score_match:
                        final_score = float(score_match.group(1))
                        detail_log = f"括号直接指定分数"
                    else:
                        bonus = 0
                        if is_red_font(cell):
                            bonus = 5
                            detail_log = f"组总分 {group_score} + 5 (红字加分)"
                        elif is_yellow_fill(cell):
                            bonus = -5
                            detail_log = f"组总分 {group_score} - 5 (黄底减分)"
                        else:
                            detail_log = f"组总分 {group_score} (无加减分)"

                        final_score = group_score + bonus

                    if final_score > 100.0:
                        detail_log += f" -> [封顶100]"
                        final_score = 100.0

                    print(f"   - {clean_name}: 最终得 {final_score} 分  [{detail_log}]")

                    if clean_name not in student_scores:
                        student_scores[clean_name] = []
                    student_scores[clean_name].append(final_score)

    print("\n" + "=" * 50)
    print("📊 跨表格汇总平均分计算完成")
    print("=" * 50)

    final_avg_scores = {}
    for name, scores_list in student_scores.items():
        if len(scores_list) > 0:
            # 使用 Decimal 进行绝对精确的求平均与两位小数四舍五入
            total = sum(Decimal(str(s)) for s in scores_list)
            avg = total / Decimal(len(scores_list))
            # 强制四舍五入到两位小数 (0.01)，并转回 float
            final_avg_scores[name] = float(avg.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP))

    target_file = '总评.xlsx'
    try:
        wb_total = load_workbook(target_file)
    except FileNotFoundError:
        print(f"\n❌ 找不到文件: {target_file}")
        return

    ws_total = wb_total.active
    headers_tot = {}
    for idx, cell in enumerate(ws_total[1]):
        if cell.value is not None:
            headers_tot[str(cell.value).strip()] = idx

    col_name = headers_tot.get('姓名')
    col_homework = headers_tot.get('大作业')

    if col_name is None or col_homework is None:
        print("\n❌ 写入失败：总评.xlsx 中未找到 '姓名' 或 '大作业' 列。")
        return

    update_count = 0
    for row in ws_total.iter_rows(min_row=2):
        raw_name = row[col_name].value
        if raw_name:
            name = str(raw_name).strip()
            if name in final_avg_scores:
                # 【核心修改点 2】：写入数据并强制规定 Excel 渲染为两位小数
                cell_homework = row[col_homework]
                cell_homework.value = final_avg_scores[name]
                cell_homework.number_format = '0.00'  # 让 Excel 严格显示为两位小数
                update_count += 1

    wb_total.save(target_file)
    print(f"\n✅ 【大作业】汇总写入完毕！已将 {update_count} 人的平均成绩(保留两位小数)写入 {target_file}")


if __name__ == "__main__":
    process_all_group_scores()